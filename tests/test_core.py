import tempfile
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace

from business_logic import calculate_due_date, compute_loan_terms, get_dashboard_metrics, get_loans_df
from data_export import create_csv_backup, create_excel_export
from db import delete_customer, delete_loan, get_connection, init_db
from import_excel import import_loans, import_repayments


class FakeWorksheet:
    """Minimal worksheet stand-in for importer tests without writing files."""

    def __init__(self, values):
        self.values = values
        self.max_row = max(row for row, _ in values) if values else 3

    def cell(self, row, column):
        return SimpleNamespace(value=self.values.get((row, column)))


class LoanManagerTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp_dir.name) / "test.db"
        init_db(self.db_path)
        self.conn = get_connection(self.db_path)
        self.conn.execute("INSERT INTO customers (customer_id, full_name) VALUES ('CUST-001', 'Test Borrower')")
        self.conn.execute("""INSERT INTO loans (loan_id, customer_id, date_taken, period, due_date, principal, rate, interest, total_due)
                          VALUES ('L001', 'CUST-001', '2026-01-01', '1 Month', '2026-02-01', 1000, 0.2, 200, 1200)""")
        self.conn.commit()

    def tearDown(self):
        self.conn.close()
        self.temp_dir.cleanup()

    def test_workbook_style_terms_use_calendar_months(self):
        self.assertEqual(calculate_due_date(date(2026, 1, 31), "1 Month"), date(2026, 2, 28))
        self.assertEqual(calculate_due_date(date(2026, 1, 1), "2 Weeks"), date(2026, 1, 15))
        self.assertEqual(compute_loan_terms(1000, 0.2), (200.0, 1200.0))

    def test_balance_and_dashboard_update_from_repayment(self):
        self.conn.execute("INSERT INTO capital (transaction_id, date, type, money_in, money_out) VALUES ('T001', '2026-01-01', 'Capital', 2000, 0)")
        self.conn.execute("""INSERT INTO repayments (payment_id, loan_id, payment_date, amount_paid)
                          VALUES ('PMT-001', 'L001', '2026-01-10', 400)""")
        self.conn.commit()
        loans = get_loans_df(self.conn)
        self.assertEqual(float(loans.iloc[0].balance), 800.0)
        metrics = get_dashboard_metrics(self.conn)
        self.assertEqual(metrics["available_cash"], 1400.0)
        self.assertEqual(metrics["outstanding_balance"], 800.0)

    def test_deleting_loan_removes_linked_repayments(self):
        self.conn.execute("""INSERT INTO repayments (payment_id, loan_id, payment_date, amount_paid)
                          VALUES ('PMT-001', 'L001', '2026-01-10', 400)""")
        self.conn.commit()
        self.assertEqual(delete_loan(self.conn, "L001"), 1)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM loans").fetchone()[0], 0)
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM repayments").fetchone()[0], 0)

    def test_customer_with_history_cannot_be_deleted(self):
        deleted, message = delete_customer(self.conn, "CUST-001")
        self.assertFalse(deleted)
        self.assertIn("loan history", message)

    def test_backup_contains_each_editable_register(self):
        import io
        import zipfile

        backup = create_csv_backup(self.conn)
        with zipfile.ZipFile(io.BytesIO(backup)) as archive:
            self.assertEqual(
                set(archive.namelist()),
                {"README.txt", "customers.csv", "loans.csv", "repayments.csv", "capital.csv"},
            )
            self.assertIn("Test Borrower", archive.read("customers.csv").decode("utf-8-sig"))

    def test_conflicting_workbook_ids_are_consolidated_without_losing_repayments(self):
        loan_sheet = FakeWorksheet({
            (4, 1): "L001", (4, 2): "Second Borrower", (4, 3): "2026-02-01", (4, 4): "1 Month",
            (4, 5): "2026-03-01", (4, 6): 2000, (4, 7): 0.1, (4, 8): 200, (4, 9): 2200,
            (4, 13): "Bank", (4, 14): "123", (4, 15): "OTHER-REF", (4, 16): "Imported",
        })
        repayment_sheet = FakeWorksheet({
            (4, 1): "PMT-001", (4, 2): "L001", (4, 4): "2026-02-10", (4, 5): 500,
            (4, 6): "Cash", (4, 7): "RECEIPT-1", (4, 8): "First payment", (4, 9): "Tester",
        })

        loans_added, loan_ids = import_loans(self.conn, loan_sheet)
        repayments_added = import_repayments(self.conn, repayment_sheet, loan_ids)
        self.conn.commit()
        self.assertEqual((loans_added, repayments_added, loan_ids["L001"]), (1, 1, "L002"))
        self.assertEqual(self.conn.execute("SELECT loan_id FROM repayments WHERE payment_id='PMT-001'").fetchone()[0], "L002")

        repeated_loans, repeated_ids = import_loans(self.conn, loan_sheet)
        repeated_payments = import_repayments(self.conn, repayment_sheet, repeated_ids)
        self.assertEqual((repeated_loans, repeated_payments, repeated_ids["L001"]), (0, 0, "L002"))

    def test_to_date_handles_various_formats(self):
        from import_excel import to_date

        self.assertEqual(to_date("10/08/2026"), "2026-08-10")
        self.assertEqual(to_date("31/08/2026"), "2026-08-31")
        self.assertEqual(to_date("2026-07-31"), "2026-07-31")
        self.assertEqual(to_date("2026-08-01 00:00:00"), "2026-08-01")
        self.assertEqual(to_date("01/09/2026"), "2026-09-01")
        self.assertEqual(to_date("1-9-2026"), "2026-09-01")
        self.assertEqual(to_date(date(2026, 8, 15)), "2026-08-15")
        self.assertIsNone(to_date(""))
        self.assertIsNone(to_date(None))
        self.assertIsNone(to_date("invalid-date"))

    def test_loan_consolidation_updates_existing_loan_without_duplicate(self):
        # In setUp, L001 is for 'Test Borrower' with principal 1000
        # When importing updated details for L001 with 'Test Borrower', it should UPDATE L001, not create L002
        updated_sheet = FakeWorksheet({
            (4, 1): "L001", (4, 2): "Test Borrower", (4, 3): "2026-01-01", (4, 4): "1 Month",
            (4, 5): "2026-02-01", (4, 6): 1000, (4, 7): 0.25, (4, 8): 250, (4, 9): 1250,
            (4, 13): "New Bank", (4, 14): "9999", (4, 15): "NEW-REF", (4, 16): "Updated notes",
        })
        loans_added, loan_ids = import_loans(self.conn, updated_sheet)
        self.conn.commit()
        self.assertEqual(loans_added, 0)
        self.assertEqual(loan_ids["L001"], "L001")
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM loans").fetchone()[0], 1)
        updated_row = self.conn.execute("SELECT * FROM loans WHERE loan_id='L001'").fetchone()
        self.assertEqual(updated_row["bank"], "New Bank")
        self.assertEqual(updated_row["notes"], "Updated notes")
        self.assertEqual(float(updated_row["rate"]), 0.25)

    def test_repayment_import_with_dd_mm_yyyy_dates(self):
        repayment_sheet = FakeWorksheet({
            (4, 1): "P001", (4, 2): "L001", (4, 4): "10/08/2026", (4, 5): 500,
            (4, 6): "Bank Transfer", (4, 7): "REF-100", (4, 8): "Partial", (4, 9): "Agent X",
        })
        added = import_repayments(self.conn, repayment_sheet, {"L001": "L001"})
        self.conn.commit()
        self.assertEqual(added, 1)
        pmt = self.conn.execute("SELECT * FROM repayments WHERE payment_id='P001'").fetchone()
        self.assertIsNotNone(pmt)
        self.assertEqual(pmt["payment_date"], "2026-08-10")
        self.assertEqual(float(pmt["amount_paid"]), 500.0)

    def test_single_excel_workbook_export_and_round_trip(self):
        import io
        import openpyxl
        from data_export import create_excel_export
        from import_excel import import_workbook

        excel_bytes = create_excel_export(self.conn)
        self.assertIsInstance(excel_bytes, bytes)
        self.assertGreater(len(excel_bytes), 1000)

        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
        self.assertEqual(
            set(wb.sheetnames),
            {"Dashboard", "Customers", "Loans", "Repayments", "Capital"},
        )
        self.assertGreater(wb["Dashboard"].max_row, 5)
        self.assertGreater(wb["Customers"].max_row, 1)
        self.assertGreater(wb["Loans"].max_row, 1)

        # Verify round-trip re-import into the same DB succeeds with zero duplicate errors
        results = import_workbook(io.BytesIO(excel_bytes), db_path=self.db_path)
        self.assertEqual(results, {"customers": 0, "capital": 0, "loans": 0, "repayments": 0})

    def test_ui_theme_styles_and_deferred_callables(self):
        import ui_theme
        self.assertIn("stAppDeployButton", ui_theme.COMMON_CSS)
        self.assertIn('[data-testid="stTab"]', ui_theme.COMMON_CSS)
        self.assertIn('aria-selected="true"', ui_theme.COMMON_CSS)
        self.assertIn('aria-expanded="false"', ui_theme.COMMON_CSS)

        # Verify deferred callable execution pattern
        def deferred_excel():
            return create_excel_export(self.conn)

        def deferred_csv():
            return create_csv_backup(self.conn)

        self.assertTrue(callable(deferred_excel))
        self.assertTrue(callable(deferred_csv))
        self.assertGreater(len(deferred_excel()), 0)
        self.assertGreater(len(deferred_csv()), 0)


if __name__ == "__main__":
    unittest.main()
