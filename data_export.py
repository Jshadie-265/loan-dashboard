"""Portable, read-only exports for a Loan Manager database."""

import csv
import io
import sqlite3
import zipfile
from datetime import datetime, timezone

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from business_logic import get_capital_df, get_dashboard_metrics, get_loans_df, get_repayments_df

EXPORT_TABLES = ("customers", "loans", "repayments", "capital")


def create_csv_backup(conn: sqlite3.Connection) -> bytes:
    """Return a ZIP containing one UTF-8 CSV file per editable register."""
    buffer = io.BytesIO()
    created_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(
            "README.txt",
            "Loan Manager data export\n\n"
            f"Created: {created_at}\n"
            "Contains the source records for Customers, Loans, Repayments, and Capital.\n"
            "Keep this archive somewhere safe as a portable backup.\n",
        )
        for table in EXPORT_TABLES:
            rows = conn.execute(f"SELECT * FROM {table} ORDER BY rowid").fetchall()
            output = io.StringIO(newline="")
            writer = csv.writer(output)
            writer.writerow([column[0] for column in conn.execute(f"SELECT * FROM {table} LIMIT 0").description])
            writer.writerows([tuple(row) for row in rows])
            archive.writestr(f"{table}.csv", output.getvalue().encode("utf-8-sig"))
    return buffer.getvalue()


def create_excel_export(conn: sqlite3.Connection) -> bytes:
    """Return a single consolidated Excel workbook (.xlsx) containing all registers as tabs."""
    wb = openpyxl.Workbook()

    header_fill = PatternFill(start_color="146356", end_color="146356", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    thin_border = Border(
        left=Side(style="thin", color="E0E0E0"),
        right=Side(style="thin", color="E0E0E0"),
        top=Side(style="thin", color="E0E0E0"),
        bottom=Side(style="thin", color="E0E0E0"),
    )

    def style_table(ws, headers: list[str], rows: list[list], currency_cols=None, date_cols=None, rate_cols=None):
        ws.views.sheetView[0].showGridLines = True
        currency_cols = set(currency_cols or [])
        date_cols = set(date_cols or [])
        rate_cols = set(rate_cols or [])

        # Write headers
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[1].height = 26

        # Write data rows
        for row_idx, row_data in enumerate(rows, 2):
            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                if col_idx in currency_cols:
                    cell.number_format = "#,##0.00"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx in rate_cols:
                    cell.number_format = "0.0%"
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif col_idx in date_cols:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[row_idx].height = 20

        # Auto-fit columns
        for col_idx in range(1, len(headers) + 1):
            max_len = len(str(headers[col_idx - 1]))
            for r in range(2, len(rows) + 2):
                cell_val = ws.cell(row=r, column=col_idx).value
                if cell_val is not None:
                    max_len = max(max_len, len(str(cell_val)))
            ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(max_len + 3, 40))

    # --- 1. Dashboard Tab ---
    ws_dash = wb.active
    ws_dash.title = "Dashboard"
    ws_dash.views.sheetView[0].showGridLines = True
    metrics = get_dashboard_metrics(conn)

    ws_dash.cell(row=1, column=1, value="LOAN MANAGER — EXECUTIVE SUMMARY").font = Font(name="Calibri", size=14, bold=True, color="146356")
    ws_dash.cell(row=2, column=1, value=f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}").font = Font(name="Calibri", size=10, italic=True, color="666666")

    kpis = [
        ("Total Capital", metrics["total_capital"], "#,##0.00"),
        ("Money Loaned", metrics["total_loaned"], "#,##0.00"),
        ("Total Repayments Collected", metrics["total_collected"], "#,##0.00"),
        ("Outstanding Balance", metrics["outstanding_balance"], "#,##0.00"),
        ("Interest Expected", metrics["total_interest_expected"], "#,##0.00"),
        ("Profit Earned", metrics["profit_earned"], "#,##0.00"),
        ("Available Cash", metrics["available_cash"], "#,##0.00"),
        ("Active Loans", metrics["active_count"], "0"),
        ("Paid Loans", metrics["paid_count"], "0"),
        ("Overdue Loans", metrics["overdue_count"], "0"),
    ]

    ws_dash.cell(row=4, column=1, value="Metric").fill = header_fill
    ws_dash.cell(row=4, column=1).font = header_font
    ws_dash.cell(row=4, column=2, value="Value (MWK / Count)").fill = header_fill
    ws_dash.cell(row=4, column=2).font = header_font
    ws_dash.row_dimensions[4].height = 24

    for idx, (label, val, fmt) in enumerate(kpis, 5):
        c1 = ws_dash.cell(row=idx, column=1, value=label)
        c1.font = bold_font
        c1.border = thin_border
        c2 = ws_dash.cell(row=idx, column=2, value=val)
        c2.font = regular_font
        c2.number_format = fmt
        c2.border = thin_border
        ws_dash.row_dimensions[idx].height = 20

    ws_dash.column_dimensions["A"].width = 32
    ws_dash.column_dimensions["B"].width = 24

    # --- 2. Customers Tab ---
    ws_cust = wb.create_sheet(title="Customers")
    cust_headers = ["Customer ID", "Full Name", "Phone Number", "National ID", "Address", "Occupation", "Collateral", "Status", "Notes"]
    cust_rows = [
        [r["customer_id"], r["full_name"], r["phone"], r["national_id"], r["address"], r["occupation"], r["collateral"], r["status"], r["notes"]]
        for r in conn.execute("SELECT customer_id, full_name, phone, national_id, address, occupation, collateral, status, notes FROM customers ORDER BY full_name").fetchall()
    ]
    style_table(ws_cust, cust_headers, cust_rows)

    # --- 3. Loans Tab ---
    ws_loans = wb.create_sheet(title="Loans")
    loan_headers = [
        "Loan ID", "Borrower", "Date Taken", "Period", "Due Date", "Principal (MWK)",
        "Rate", "Interest (MWK)", "Total Due (MWK)", "Amount Paid (MWK)", "Balance (MWK)",
        "Status", "Bank", "Account", "Reference", "Notes"
    ]
    loans_df = get_loans_df(conn)
    loan_rows = []
    if not loans_df.empty:
        for _, row in loans_df.iterrows():
            loan_rows.append([
                row["loan_id"], row["borrower"], row["date_taken"], row["period"], row["due_date"],
                float(row["principal"]), float(row["rate"]), float(row["interest"]), float(row["total_due"]),
                float(row["amount_paid"]), float(row["balance"]), row["status"],
                row["bank"], row["account"], row["reference"], row["notes"]
            ])
    style_table(ws_loans, loan_headers, loan_rows, currency_cols={6, 8, 9, 10, 11}, date_cols={3, 5}, rate_cols={7})

    # --- 4. Repayments Tab ---
    ws_rep = wb.create_sheet(title="Repayments")
    rep_headers = ["Payment ID", "Loan ID", "Borrower", "Payment Date", "Amount Paid (MWK)", "Payment Method", "Reference", "Notes", "Recorded By"]
    rep_df = get_repayments_df(conn)
    rep_rows = []
    if not rep_df.empty:
        for _, row in rep_df.iterrows():
            rep_rows.append([
                row["payment_id"], row["loan_id"], row["borrower"], row["payment_date"],
                float(row["amount_paid"]), row["payment_method"], row["reference"], row["notes"], row["recorded_by"]
            ])
    style_table(ws_rep, rep_headers, rep_rows, currency_cols={5}, date_cols={4})

    # --- 5. Capital Tab ---
    ws_cap = wb.create_sheet(title="Capital")
    cap_headers = ["Transaction ID", "Date", "Type", "Description", "Money In (MWK)", "Money Out (MWK)", "Running Balance (MWK)"]
    cap_df = get_capital_df(conn)
    cap_rows = []
    if not cap_df.empty:
        for _, row in cap_df.iterrows():
            cap_rows.append([
                row["transaction_id"], row["date"], row["type"], row["description"],
                float(row["money_in"]), float(row["money_out"]), float(row["running_balance"])
            ])
    style_table(ws_cap, cap_headers, cap_rows, currency_cols={5, 6, 7}, date_cols={2})

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
